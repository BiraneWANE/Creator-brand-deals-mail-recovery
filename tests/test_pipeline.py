"""Test de bout en bout sur la boîte fictive : chaque piège connu doit être traité comme prévu."""

import sqlite3
from datetime import date

from openpyxl import load_workbook

from branddeals.classify import RuleReader
from branddeals.config import Rules
from branddeals.pipeline import LABEL, run
from branddeals.sources.synthetic import BRANDS, ME, SyntheticMailbox


def _run(tmp_path):
    rules = Rules.load()
    rules.own_addresses = {ME}
    box = SyntheticMailbox(seed=7)
    return box, run(box, RuleReader(rules), rules, tmp_path, today=date(2026, 10, 27), workers=2, batch_size=3)


def test_end_to_end(tmp_path):
    box, res = _run(tmp_path)
    s = res.stats
    # 12 marques fictives : 9 vraies collabs (payées, offertes, campagnes), 3 refus polis
    assert s["partenariats_retenus"] == 9
    assert s["interlocuteurs_ecartes_par_regles"] >= 3
    # 3 collabs payées : 1 contrat + 1 facture chacune, logos de signature ignorés
    assert s["documents_extraits"] == 6
    # 3 briefs trop lourds : listés pour récupération manuelle, pas perdus
    assert s["documents_a_recuperer_a_la_main"] == 3
    # le piège « numéro de facture réutilisé » est détecté
    assert any(a.sujet == "numéro de facture réutilisé" for a in res.alerts)
    # seule écriture dans la boîte : le libellé, et uniquement sur les conversations retenues
    assert sum(LABEL in v for v in box.labels.values()) == s["conversations_libellees"] == 9


def test_outputs_are_consistent(tmp_path):
    _, res = _run(tmp_path)
    con = sqlite3.connect(res.outputs["base"])
    assert con.execute("SELECT COUNT(*) FROM partenariat").fetchone()[0] == 9
    assert con.execute("SELECT COUNT(*) FROM v_numeros_factures_reutilises").fetchone()[0] == 1
    # intégrité : chaque document pointe vers une conversation connue
    orphans = con.execute("SELECT COUNT(*) FROM document d LEFT JOIN conversation c ON c.id = d.id_conversation "
                          "WHERE c.id IS NULL").fetchone()[0]
    assert orphans == 0
    wb = load_workbook(res.outputs["classeur"])
    assert wb.sheetnames == ["Synthèse", "Partenariats", "Factures émises", "Contrats et devis",
                             "Briefs et autres", "À télécharger à la main", "Alertes"]
    assert wb["Partenariats"].max_row == 10


def test_synthetic_mailbox_only_uses_reserved_domains():
    box = SyntheticMailbox()
    addresses = {a for t in box.threads.values() for a in t.addresses()}
    assert all(a.endswith("example.com") for a in addresses)
    assert len(BRANDS) == 12
