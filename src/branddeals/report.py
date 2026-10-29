"""Étape 6 : le livrable métier, un classeur Excel lisible sans aucune compétence technique.

Onglets : Synthèse · Partenariats (relances surlignées) · Factures émises · Contrats et devis ·
Briefs et autres · À télécharger à la main · Alertes. Les totaux sont des formules Excel,
pas des valeurs figées : le classeur reste juste si la créatrice corrige une ligne.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import DocType, Document, Partnership
from .quality import Alerte

HEAD = PatternFill("solid", fgColor="1F4E78")
RELANCE = PatternFill("solid", fgColor="FFF2CC")
ALERTE = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
FONT = "Arial"


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes, ws.auto_filter.ref = "B2", ws.dimensions
    return ws


def write_workbook(path: str | Path, partnerships: list[Partnership], docs: list[Document],
                   alerts: list[Alerte], stats: dict[str, int]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    ps = sorted(partnerships, key=lambda p: p.dernier_echange, reverse=True)

    synth = wb.create_sheet("Synthèse")
    synth["A1"] = "Historique des partenariats"
    synth["A1"].font = Font(name=FONT, bold=True, size=14)
    n = len(ps) + 1
    lines = [("Partenariats retenus", f"=COUNTA(Partenariats!B2:B{n})"),
             ("À relancer", f'=COUNTIF(Partenariats!A2:A{n},"OUI")'),
             ("Collabs payées", f'=COUNTIF(Partenariats!K2:K{n},"payée")'),
             ("Factures en attente de paiement", f'=COUNTIF(Partenariats!K2:K{n},"en attente")')]
    lines += [(k.replace("_", " ").capitalize(), v) for k, v in stats.items()]
    for i, (label, value) in enumerate(lines, start=3):
        synth.cell(i, 1, label).font = Font(name=FONT)
        synth.cell(i, 2, value).font = Font(name=FONT, bold=True)
    synth.column_dimensions["A"].width, synth.column_dimensions["B"].width = 42, 14
    synth.cell(len(lines) + 4, 1, "Lignes jaunes : collab réussie ou contact chaleureux, sans échange depuis 6 mois.").font = \
        Font(name=FONT, italic=True, size=9)

    ws = _sheet(wb, "Partenariats",
                ["À relancer", "Marque / agence", "Contact", "Email", "Premier échange", "Dernier échange", "Type",
                 "Montant (€)", "N° facture", "Remarque", "Statut", "Lien"],
                [["OUI" if p.a_relancer else "", p.marque, p.contact_nom, p.contact_email, p.premier_echange,
                  p.dernier_echange, p.type.value, p.montant, p.facture, p.remarque, p.statut.value, p.lien]
                 for p in ps],
                [10, 28, 20, 28, 13, 13, 15, 11, 12, 60, 12, 10])
    for row in ws.iter_rows(min_row=2):
        row[4].number_format = row[5].number_format = "DD/MM/YYYY"
        row[7].number_format = "#,##0 €"
        if row[0].value == "OUI":
            for c in row:
                c.fill = RELANCE

    inv = [d for d in docs if d.doc_type is DocType.FACTURE_EMISE]
    reused = {a.detail.split()[1] for a in alerts if a.sujet == "numéro de facture réutilisé"}
    wf = _sheet(wb, "Factures émises", ["Date", "N° facture", "Client", "Montant (€)", "Fichier", "Remarque"],
                [[d.sent_at, d.numero, d.counterparty_key, d.montant, d.path, d.remarque] for d in inv],
                [12, 12, 30, 12, 55, 50])
    for row in wf.iter_rows(min_row=2):
        row[0].number_format, row[3].number_format = "DD/MM/YYYY", "#,##0 €"
        if row[1].value in reused:
            for c in row:
                c.fill = ALERTE
    last = len(inv) + 1
    wf.cell(last + 2, 3, "Total des factures retrouvées").font = Font(name=FONT, bold=True)
    total = wf.cell(last + 2, 4, f"=SUM(D2:D{last})")
    total.font, total.number_format = Font(name=FONT, bold=True), "#,##0 €"

    contracts = (DocType.CONTRAT, DocType.DEVIS, DocType.BON_DE_COMMANDE)
    _sheet(wb, "Contrats et devis", ["Date", "Type", "Interlocuteur", "Fichier"],
           [[d.sent_at, d.doc_type.value, d.counterparty_key, d.path] for d in docs if d.doc_type in contracts and d.path],
           [12, 16, 30, 60])
    _sheet(wb, "Briefs et autres", ["Date", "Type", "Interlocuteur", "Fichier"],
           [[d.sent_at, d.doc_type.value, d.counterparty_key, d.path] for d in docs
            if d.doc_type not in (*contracts, DocType.FACTURE_EMISE) and d.path],
           [12, 14, 30, 60])
    _sheet(wb, "À télécharger à la main", ["Date", "Interlocuteur", "Pièces jointes", "Pourquoi"],
           [[d.sent_at, d.counterparty_key, d.filename, d.remarque] for d in docs if d.path is None],
           [12, 30, 50, 45])
    _sheet(wb, "Alertes", ["Niveau", "Sujet", "Détail"], [[a.niveau, a.sujet, a.detail] for a in alerts], [12, 30, 80])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
